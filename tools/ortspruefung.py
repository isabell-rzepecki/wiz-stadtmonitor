"""Ortsprüfung: unklare und mehrdeutige Ortsbezüge zur Prüfung durch Ortskundige ausgeben und geprüft zurückspielen.

Zweck: Die Verortung aus Titeln (pipeline/verortung.py, Stufe B) kann Straßennamen, die es in mehreren Ortsteilen gibt,
nicht immer eindeutig zuordnen, und Namen von Einrichtungen (Freibad, Kesperschule) stehen in keinem Straßenverzeichnis.
Diese Fälle gehen als Excel-Tabelle an Ortskundige; die ausgefüllte Tabelle wird eingelesen und als
config/verortung_geprueft.tsv versioniert. Der Seitenbau wendet sie an (verortung.geprueft_anwenden).

Aufruf aus dem Wurzelverzeichnis des Repositorys, nach `python -m pipeline build`:
  python tools/ortspruefung.py erstellen ZIEL.xlsx     Tabelle mit den offenen Fällen (bereits geprüfte fehlen)
  python tools/ortspruefung.py lage ZIEL.xlsx [--ortsteil Kernstadt]
                                                        Nachtrag: geprüfte Orte ohne Straße, mit Auswahlliste der Straßen
  python tools/ortspruefung.py einlesen DATEI.xlsx      ausgefüllte Tabelle prüfen und in config/verortung_geprueft.tsv übernehmen
Die Excel-Datei selbst gehört nicht ins Repository; versioniert wird nur die geprüfte TSV-Datei (ohne Namen).
"""
import argparse
import csv
import json
import re
import sys
from datetime import date
from pathlib import Path

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))
from pipeline import config, parse, verortung  # noqa: E402

AUDIT = config.EXPORT / "audit" / "verortung_stufe_b.tsv"
FAELLE = {
    "offen": "Straßenname in mehreren Ortsteilen, nicht zugeordnet",
    "vorrang": "Straßenname in mehreren Ortsteilen, über die Vorrangregel zugeordnet",
    "eingegrenzt": "Straßenname in mehreren Ortsteilen, über Gremium oder Titel eingegrenzt",
    "ortsname": "Name einer Einrichtung oder Lage, nicht im Straßenverzeichnis",
    "lage": "Ortsteil geprüft, Straße fehlt",
}
# Namen von Einrichtungen und Lagen, die das Straßenverzeichnis nicht kennt; nur als Kandidaten für die Prüfung
ORTSNAME = re.compile(r"\b(?:[A-ZÄÖÜ][\wäöüß-]*(?:bad|schule|hof|halle|haus|platz|brücke|gewächshaus|park|friedhof|garten)|"
                      r"Bahnhof|Marktplatz|Freibad|Hallenbad|Kindergarten|Kita)\b")
SPALTEN = [  # (Feld, Überschrift, Breite, auszufüllen)
    ("nr", "Nr.", 5, False), ("vorgang", "Vorgang", 16, False), ("fall", "Fall", 30, False), ("titel", "Titel des Tagesordnungspunkts", 60, False),
    ("gremium", "Gremium", 26, False), ("datum", "Sitzung am", 11, False), ("treffer", "Fundstelle im Titel", 22, False),
    ("bisher", "Bisherige Zuordnung", 24, False), ("kandidaten", "Mögliche Ortsteile laut Verzeichnis", 26, False),
    ("quelle", "Einladung (PDF)", 18, False),
    ("ergebnis", "Ergebnis", 16, True), ("ortsteil", "Ortsteil", 18, True), ("strasse", "Straße oder Lage (ohne Hausnummer)", 26, True),
    ("hinweis", "Hinweis (ohne Namen von Personen)", 30, True), ("geprueft_am", "Geprüft am (TT.MM.JJJJ)", 14, True),
]
ANLEITUNG = [
    "Ortsprüfung: Bitte für jede Zeile prüfen, welchen Ort der Tagesordnungspunkt meint.",
    "",
    "Die gelben Spalten ausfüllen:",
    "• Ergebnis: „bestätigt“ (bisherige Zuordnung stimmt), „korrigiert“ (anderer oder fehlender Ort, dann Ortsteil angeben),",
    "  „kein Ortsbezug“ (der Punkt betrifft keinen bestimmten Ort, z. B. ein Bericht über einen Betrieb) oder „unklar“.",
    "• Ortsteil: aus der Auswahlliste.",
    "• Straße oder Lage: Straßenname, an dem der Ort liegt, ohne Hausnummer (z. B. „Am Sande“). Leer lassen, wenn nur der Ortsteil bekannt ist.",
    "• Hinweis: kurze Begründung oder Quelle (z. B. „Einladung S. 3“). Bitte keine Namen von Personen, keine Telefonnummern.",
    "• Geprüft am: Datum der Prüfung.",
    "",
    "Die grauen Spalten bitte nicht ändern; Vorgang und Fundstelle verbinden die Zeile mit dem Stadtmonitor.",
    "Quelle der Titel: Amtliche Bekanntmachungen der Stadt (Einladungen). Die ausgefüllte Datei an die verantwortliche Stelle zurückgeben.",
]
ANLEITUNG_LAGE = [
    "Ortsprüfung, Nachtrag: An welcher Straße liegt die Einrichtung?",
    "",
    "Diese Orte sind bereits geprüft, aber nur einem Ortsteil zugeordnet. Mit der Straße erscheinen sie als Punkt auf der Karte.",
    "• Straße oder Lage: aus der Auswahlliste wählen (amtliche Schreibweise), ohne Hausnummer. Leer lassen, wenn unbekannt.",
    "• Ergebnis und Ortsteil sind vorausgefüllt; nur ändern, wenn sie nicht stimmen. „kein Ortsbezug“, wenn der Punkt keinen Ort meint.",
    "• Hinweis: kurze Quelle, ohne Namen von Personen. Geprüft am: Datum der Prüfung.",
    "",
    "Die grauen Spalten bitte nicht ändern; Vorgang und Fundstelle verbinden die Zeile mit dem Stadtmonitor.",
    "Die ausgefüllte Datei an die verantwortliche Stelle zurückgeben.",
]


def lage_faelle(vorgaenge, geprueft, ortsteil=None):
    """Nachtrag: als „korrigiert“ geprüfte Orte ohne Straße, Ergebnis und Ortsteil vorausgefüllt."""
    v_nach = {v["reference"]: v for v in vorgaenge}
    out = []
    for zs in geprueft.values():
        for g in zs:
            v = v_nach.get(g["vorgang"])
            if not v or g["ergebnis"] != "korrigiert" or g["strasse"] or (ortsteil and g["ortsteil"] != ortsteil):
                continue
            p = v["process"][0]
            out.append({"vorgang": v["reference"], "fall": FAELLE["lage"], "titel": parse.ohne_namen(v["name"]), "gremium": p["gremium"],
                        "datum": p["date"], "treffer": g["treffer"], "bisher": f"{g['ortsteil']} (geprüft am {g['geprueft_am']})",
                        "kandidaten": "", "quelle": p["pdf"], "ergebnis": "korrigiert", "ortsteil": g["ortsteil"], "hinweis": g["hinweis"]})
    out.sort(key=lambda z: (z["treffer"], z["datum"]))
    for i, z in enumerate(out, 1):
        z["nr"] = i
    return out


def strassen(ortsteile):
    """Straßen- und Lagenamen des Verzeichnisses in den Gemarkungen dieser Ortsteile (nur solche, die die Verortung kennt)."""
    gks = {(verortung._ortsteil(name=o) or {}).get("gemarkung_nr") for o in ortsteile}
    return sorted(n for n, je in verz().namen.items() if gks & set(je))


def faelle(vorgaenge, audit, geprueft):
    """Zeilen für die Prüfung: je Vorgang und Fundstelle eine, ohne bereits geprüfte (außer „unklar“)."""
    erledigt = {(z["vorgang"], z["treffer"]) for zs in geprueft.values() for z in zs if z["ergebnis"] != "unklar"}
    v_nach = {v["reference"]: v for v in vorgaenge}
    out, gesehen = [], set()

    def zeile(v, fall, treffer, bisher, kandidaten):
        if (v["reference"], treffer) in gesehen or (v["reference"], treffer) in erledigt:
            return
        gesehen.add((v["reference"], treffer))
        p = v["process"][0]
        out.append({"vorgang": v["reference"], "fall": FAELLE[fall], "titel": parse.ohne_namen(v["name"]), "gremium": p["gremium"],
                    "datum": p["date"], "treffer": treffer, "bisher": bisher, "kandidaten": kandidaten, "quelle": p["pdf"]})

    for a in audit:
        v = v_nach.get(a["vorgang"])
        if not v or a["mehrdeutig"] != "ja" or a["methode_detail"] != "B-strasse":
            continue
        loc = next((l for l in v["locations"] if l["name"] == a["treffer"] and l["level"] == "strasse"), None)
        if not a["konfidenz"]:
            fall = "offen"
        elif loc and loc["method"] == "strasse+vorrang":
            fall = "vorrang"
        else:
            fall = "eingegrenzt"
        kand = sorted(filter(None, ((verortung._ortsteil(gk=gk) or {}).get("name") for gk in verz().namen.get(a["treffer"], {}))))
        zeile(v, fall, a["treffer"], f"{a['treffer']} ({loc['ortsteil']})" if loc else "", ", ".join(kand))
    for v in vorgaenge:
        # nur ganz ohne Ort; mit Ortsbeirat ist wenigstens der Ortsteil bekannt
        if v["verfahrenspunkt"] or v["locations"]:
            continue
        for m in ORTSNAME.finditer(v["name"]):
            if m.group(0) not in verz().namen:
                zeile(v, "ortsname", m.group(0), "", "")
    reihen = list(FAELLE)
    out.sort(key=lambda z: (reihen.index(next(k for k, t in FAELLE.items() if t == z["fall"])), z["treffer"], z["datum"]))
    for i, z in enumerate(out, 1):
        z["nr"] = i
    return out


_VERZ = []


def verz():
    if not _VERZ:
        _VERZ.append(verortung.Verzeichnis())
    return _VERZ[0]


def erstellen(ziel, vorgaenge=None, audit=None, geprueft=None, lage=False, ortsteil=None):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.worksheet.datavalidation import DataValidation

    if vorgaenge is None:
        vorgaenge = json.loads(config.VORGAENGE.read_text(encoding="utf-8"))["vorgaenge"]
    if audit is None and not lage:
        with AUDIT.open(encoding="utf-8", newline="") as f:
            audit = list(csv.DictReader(f, delimiter="\t"))
    geprueft = verortung.geprueft_laden() if geprueft is None else geprueft
    zeilen = lage_faelle(vorgaenge, geprueft, ortsteil) if lage else faelle(vorgaenge, audit, geprueft)
    wb = Workbook()
    ws = wb.active
    ws.title = "Prüfung"
    grau, gelb = PatternFill("solid", fgColor="E7E6E6"), PatternFill("solid", fgColor="FFF2CC")
    for c, (_, kopf, breite, aus) in enumerate(SPALTEN, 1):
        zelle = ws.cell(row=1, column=c, value=kopf)
        zelle.font, zelle.fill = Font(bold=True), gelb if aus else grau
        zelle.alignment = Alignment(wrap_text=True, vertical="top")
        ws.column_dimensions[zelle.column_letter].width = breite
    for r, z in enumerate(zeilen, 2):
        for c, (feld, _, _, aus) in enumerate(SPALTEN, 1):
            zelle = ws.cell(row=r, column=c, value=z.get(feld, ""))
            zelle.alignment = Alignment(wrap_text=True, vertical="top")
            if aus:
                zelle.fill = gelb
            if feld == "quelle" and z["quelle"]:
                zelle.value, zelle.hyperlink, zelle.font = "Einladung öffnen", z["quelle"], Font(color="0563C1", underline="single")
    ws.freeze_panes = "B2"
    ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(SPALTEN)).column_letter}{max(len(zeilen) + 1, 2)}"
    listen = wb.create_sheet("Listen")
    for i, e in enumerate(verortung.ERGEBNISSE, 1):
        listen.cell(row=i, column=1, value=e)
    ortsteile = [o["name"] for o in config.STADT["ortsteile"]]
    for i, o in enumerate(ortsteile, 1):
        listen.cell(row=i, column=2, value=o)
    listen.sheet_state = "hidden"
    letzte = len(zeilen) + 1
    spalte = {f: ws.cell(row=1, column=i).column_letter for i, (f, *_rest) in enumerate(SPALTEN, 1)}
    for feld, formel in (("ergebnis", f"=Listen!$A$1:$A${len(verortung.ERGEBNISSE)}"), ("ortsteil", f"=Listen!$B$1:$B${len(ortsteile)}")):
        dv = DataValidation(type="list", formula1=formel, allow_blank=True, showErrorMessage=True)
        ws.add_data_validation(dv)
        dv.add(f"{spalte[feld]}2:{spalte[feld]}{letzte}")
    if lage:  # Auswahlliste der Straßen; andere Eingaben nur mit Warnung, einlesen prüft gegen das Verzeichnis
        namen = strassen({z["ortsteil"] for z in zeilen})
        for i, n in enumerate(namen, 1):
            listen.cell(row=i, column=3, value=n)
        dv = DataValidation(type="list", formula1=f"=Listen!$C$1:$C${max(len(namen), 1)}", allow_blank=True,
                            showErrorMessage=True, errorStyle="warning",
                            error="Bitte die amtliche Schreibweise aus der Liste wählen.", errorTitle="Straße nicht in der Liste")
        ws.add_data_validation(dv)
        dv.add(f"{spalte['strasse']}2:{spalte['strasse']}{letzte}")
    hilfe = wb.create_sheet("Anleitung", 0)
    for i, text in enumerate(ANLEITUNG_LAGE if lage else ANLEITUNG, 1):
        hilfe.cell(row=i, column=1, value=text)
    hilfe.column_dimensions["A"].width = 130
    hilfe["A1"].font = Font(bold=True, size=13)
    wb.active = 0
    Path(ziel).parent.mkdir(parents=True, exist_ok=True)
    wb.save(ziel)
    return zeilen


def _datum(wert):
    if hasattr(wert, "strftime"):
        return wert.strftime("%Y-%m-%d")
    w = str(wert or "").strip()
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", w)
    if m:
        return date(int(m.group(3)), int(m.group(2)), int(m.group(1))).isoformat()
    if re.fullmatch(r"\d{4}-\d{2}-\d{2}", w):
        return w
    return None


# Freitext ohne Personennamen, Kontaktdaten und Hausnummern (wie in der Erfassung)
VERBOTEN = [(parse.PERSON, "Name einer Person"), (re.compile(r"@|\b\d[\d /-]{5,}\d\b"), "Kontaktdaten")]
HAUSNUMMER = re.compile(r"\d")


def einlesen(datei, ziel=verortung.GEPRUEFT, heute=None):
    """Ausgefüllte Tabelle prüfen und in die TSV übernehmen. Rückgabe (übernommen, fehler); bei Fehlern wird nichts geschrieben."""
    from openpyxl import load_workbook

    ws = load_workbook(datei, read_only=True, data_only=True)["Prüfung"]
    reihen = ws.iter_rows(values_only=True)
    kopf = list(next(reihen))
    idx = {feld: kopf.index(k) for feld, k, *_ in SPALTEN}
    ortsteile = {o["name"] for o in config.STADT["ortsteile"]}
    neu, fehler = [], []
    for n, r in enumerate(reihen, 2):
        z = {f: ("" if r[i] is None else r[i]) for f, i in idx.items()}
        erg = str(z["ergebnis"]).strip()
        if not erg:
            continue
        zeile = {"vorgang": str(z["vorgang"]).strip(), "treffer": str(z["treffer"]).strip(), "ergebnis": erg,
                 "ortsteil": str(z["ortsteil"]).strip(), "strasse": re.sub(r"\s+", " ", str(z["strasse"])).strip(),
                 "hinweis": re.sub(r"\s+", " ", str(z["hinweis"])).strip(), "geprueft_am": _datum(z["geprueft_am"]) or (heute or date.today()).isoformat()}
        probleme = []
        if erg not in verortung.ERGEBNISSE:
            probleme.append(f"Ergebnis „{erg}“ unbekannt")
        if not zeile["vorgang"] or not zeile["treffer"]:
            probleme.append("Vorgang oder Fundstelle fehlt")
        if erg == "korrigiert" and zeile["ortsteil"] not in ortsteile:
            probleme.append("bei „korrigiert“ einen Ortsteil aus der Liste angeben")
        if zeile["ortsteil"] and zeile["ortsteil"] not in ortsteile:
            probleme.append(f"Ortsteil „{zeile['ortsteil']}“ unbekannt")
        if erg == "bestätigt" and not z["bisher"]:
            probleme.append("„bestätigt“ ohne bisherige Zuordnung; bitte „korrigiert“ mit Ortsteil")
        for feld in ("strasse", "hinweis"):
            for muster, was in VERBOTEN:
                if muster.search(zeile[feld]):
                    probleme.append(f"{feld}: {was} entfernen")
        if HAUSNUMMER.search(zeile["strasse"]):
            probleme.append("strasse: ohne Hausnummer")
        elif zeile["strasse"] and zeile["ortsteil"] in ortsteile and zeile["strasse"] not in strassen({zeile["ortsteil"]}):
            probleme.append(f"Straße „{zeile['strasse']}“ steht nicht im Verzeichnis von {zeile['ortsteil']}; Schreibweise aus der Auswahlliste")
        if probleme:
            fehler.append(f"Zeile {n} ({zeile['vorgang']}, {zeile['treffer']}): " + "; ".join(probleme))
        else:
            neu.append(zeile)
    if fehler:
        return [], fehler
    alt = [z for zs in verortung.geprueft_laden(ziel).values() for z in zs]
    schluessel = {(z["vorgang"], z["treffer"]) for z in neu}
    alle = sorted([z for z in alt if (z["vorgang"], z["treffer"]) not in schluessel] + neu, key=lambda z: (z["vorgang"], z["treffer"]))
    with Path(ziel).open("w", encoding="utf-8", newline="") as f:
        w = csv.DictWriter(f, fieldnames=verortung.GEPRUEFT_FELDER, delimiter="\t", lineterminator="\n")
        w.writeheader()
        w.writerows(alle)
    return neu, []


def main(argv=None):
    ap = argparse.ArgumentParser(description=__doc__.split("\n\n")[0])
    sub = ap.add_subparsers(dest="befehl", required=True)
    sub.add_parser("erstellen").add_argument("ziel")
    sub.add_parser("einlesen").add_argument("datei")
    lp = sub.add_parser("lage")
    lp.add_argument("ziel")
    lp.add_argument("--ortsteil")
    a = ap.parse_args(argv)
    if a.befehl in ("erstellen", "lage"):
        zeilen = erstellen(a.ziel, lage=a.befehl == "lage", ortsteil=getattr(a, "ortsteil", None))
        print(f"{len(zeilen)} Fälle in {a.ziel}")
        return 0
    neu, fehler = einlesen(a.datei)
    for f in fehler:
        print(f, file=sys.stderr)
    if fehler:
        print("Nichts übernommen; bitte die Zeilen korrigieren.", file=sys.stderr)
        return 1
    print(f"{len(neu)} geprüfte Zeilen in {verortung.GEPRUEFT.relative_to(ROOT)} übernommen; danach python -m pipeline build")
    return 0


if __name__ == "__main__":
    sys.exit(main())
