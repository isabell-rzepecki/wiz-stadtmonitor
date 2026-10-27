"""Ortsprüfung: Tabelle erstellen, ausgefüllt einlesen, im Seitenbau anwenden."""
import json
import sys
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

from pipeline import config, verortung

sys.path.insert(0, str(config.ROOT / "tools"))
import ortspruefung  # noqa: E402

NAMEN = {
    "Ermschwerder Straße (Witzenhausen)": {"2246": {"n": 40, "lon": 9.86, "lat": 51.35}},
    "Ermschwerder Straße (Blickershausen)": {"2125": {"n": 5, "lon": 9.79, "lat": 51.37}},
    "Bachstraße (Witzenhausen)": {"2246": {"n": 9, "lon": 9.85, "lat": 51.35}},
    "Bachstraße (Hundelshausen)": {"2173": {"n": 9, "lon": 9.85, "lat": 51.29}},
    "Am Sande": {"2246": {"n": 12, "lon": 9.858, "lat": 51.344}},
}
PDF = "https://example.org/e.pdf"


def vorgang(ref, titel, locations=(), gremium="Bau- und Verkehrsausschuss"):
    return {"reference": ref, "name": titel, "verfahrenspunkt": False, "locations": list(locations),
            "process": [{"gremium": gremium, "date": "2026-05-01", "pdf": PDF}]}


class Ortspruefung(unittest.TestCase):
    def setUp(self):
        self.tmp = tempfile.TemporaryDirectory()
        self.dir = Path(self.tmp.name)
        p = self.dir / "v.json"
        p.write_text(json.dumps({"namen": NAMEN}), encoding="utf-8")
        self.verz = verortung.Verzeichnis(p)
        ortspruefung._VERZ[:] = [self.verz]
        erm = {"level": "strasse", "name": "Ermschwerder Straße", "ortsteil": "Kernstadt", "stage": "B", "method": "strasse+vorrang",
               "confidence": "mittel", "evidence": "Ermschwerder Straße", "lon": 9.86, "lat": 51.35}
        self.vorgaenge = [
            vorgang("BV/013/2025", "Vorstellung Entwurfsplanung Ermschwerder Straße", [erm]),
            vorgang("147/2025-TOP3.2", "Bachlauf in der unteren Bachstraße",
                    [{"level": "ortsteil", "name": "Roßbach", "ortsteil": "Roßbach", "stage": "A", "method": "gremium_ortsbeirat"}], "Ortsbeirat Roßbach"),
            vorgang("STV/001/2026", "Antrag zum Freibad"),
            vorgang("HFA-1", "Bericht Baubetriebshof, Gespräch mit Erika Muster"),
        ]
        self.audit = [
            {"vorgang": "BV/013/2025", "mehrdeutig": "ja", "methode_detail": "B-strasse", "treffer": "Ermschwerder Straße", "konfidenz": "mittel"},
            {"vorgang": "147/2025-TOP3.2", "mehrdeutig": "ja", "methode_detail": "B-strasse", "treffer": "Bachstraße", "konfidenz": ""},
        ]

    def tearDown(self):
        ortspruefung._VERZ.clear()
        self.tmp.cleanup()

    def tabelle(self):
        ziel = self.dir / "pruefung.xlsx"
        zeilen = ortspruefung.erstellen(ziel, self.vorgaenge, self.audit, geprueft={})
        return ziel, zeilen

    def ausfuellen(self, ziel, werte):
        wb = load_workbook(ziel)
        ws = wb["Prüfung"]
        kopf = [c.value for c in ws[1]]
        spalte = {f: kopf.index(k) + 1 for f, k, *_ in ortspruefung.SPALTEN}
        for r in range(2, ws.max_row + 1):
            treffer = ws.cell(row=r, column=spalte["treffer"]).value
            for feld, wert in werte.get(treffer, {}).items():
                ws.cell(row=r, column=spalte[feld], value=wert)
        wb.save(ziel)

    def test_faelle_und_tabelle(self):
        ziel, zeilen = self.tabelle()
        self.assertEqual([(z["treffer"], z["fall"]) for z in zeilen], [
            ("Bachstraße", ortspruefung.FAELLE["offen"]), ("Ermschwerder Straße", ortspruefung.FAELLE["vorrang"]),
            ("Baubetriebshof", ortspruefung.FAELLE["ortsname"]), ("Freibad", ortspruefung.FAELLE["ortsname"])])
        self.assertEqual(zeilen[0]["kandidaten"], "Hundelshausen, Kernstadt")
        wb = load_workbook(ziel)
        self.assertEqual(wb.sheetnames[:2], ["Anleitung", "Prüfung"])
        text = " ".join(str(c.value) for row in wb["Prüfung"].iter_rows() for c in row)
        self.assertNotIn("Erika", text)  # Namen ohne Anrede werden ersetzt
        self.assertEqual(len(wb["Prüfung"].data_validations.dataValidation), 2)

    def test_einlesen_und_anwenden(self):
        ziel, _ = self.tabelle()
        self.ausfuellen(ziel, {
            "Ermschwerder Straße": {"ergebnis": "bestätigt", "geprueft_am": "02.11.2026"},
            "Bachstraße": {"ergebnis": "korrigiert", "ortsteil": "Roßbach", "hinweis": "Bachlauf am Ortsrand, Einladung S. 2"},
            "Freibad": {"ergebnis": "korrigiert", "ortsteil": "Kernstadt", "strasse": "Am Sande"},
            "Baubetriebshof": {"ergebnis": "kein Ortsbezug"},
        })
        tsv = self.dir / "geprueft.tsv"
        neu, fehler = ortspruefung.einlesen(ziel, tsv)
        self.assertEqual(fehler, [])
        self.assertEqual(len(neu), 4)
        g = verortung.geprueft_laden(tsv)
        self.assertEqual(g["BV/013/2025"][0]["geprueft_am"], "2026-11-02")
        v = {x["reference"]: x for x in self.vorgaenge}
        locs = verortung.geprueft_anwenden(v["BV/013/2025"]["locations"], g["BV/013/2025"], self.verz, PDF)
        self.assertEqual([(l["name"], l["confidence"]) for l in locs], [("Ermschwerder Straße", "hoch")])
        locs = verortung.geprueft_anwenden(v["STV/001/2026"]["locations"], g["STV/001/2026"], self.verz, PDF)
        self.assertEqual([(l["level"], l["name"], l["method"], l["lon"]) for l in locs], [("strasse", "Am Sande", "geprueft", 9.858)])
        # Bachstraße steht in Roßbach nicht im Verzeichnis: Ortsteil statt Punkt; der Gremiumsort (Stufe A) bleibt
        locs = verortung.geprueft_anwenden(v["147/2025-TOP3.2"]["locations"], g["147/2025-TOP3.2"], self.verz, PDF)
        self.assertEqual([(l["stage"], l["level"], l["name"]) for l in locs], [("A", "ortsteil", "Roßbach"), ("B", "ortsteil", "Roßbach")])
        # erneut erstellen: geprüfte Fälle fehlen
        self.assertEqual(ortspruefung.faelle(self.vorgaenge, self.audit, g), [])

    def test_einlesen_lehnt_ab_und_schreibt_nichts(self):
        ziel, _ = self.tabelle()
        self.ausfuellen(ziel, {
            "Freibad": {"ergebnis": "korrigiert", "strasse": "Am Sande 3"},
            "Bachstraße": {"ergebnis": "unklar", "hinweis": "laut Frau Muster, Tel. 05542 123456"},
            "Baubetriebshof": {"ergebnis": "bestätigt"},
        })
        tsv = self.dir / "geprueft.tsv"
        neu, fehler = ortspruefung.einlesen(ziel, tsv)
        self.assertEqual(neu, [])
        self.assertFalse(tsv.exists())
        alles = " ".join(fehler)
        for teil in ("Ortsteil", "Hausnummer", "Name einer Person", "Kontaktdaten", "„bestätigt“ ohne bisherige Zuordnung"):
            self.assertIn(teil, alles)


if __name__ == "__main__":
    unittest.main()
